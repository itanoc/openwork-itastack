#!/usr/bin/env python3
"""write_scope.py — Write task plan data to a copy of the Sales Scope Template.

Reads a JSON task plan from stdin or a file and populates the Tasklist tab:
  - Column A: Task descriptions (organized by day)
  - Column C: Time estimates (minimum hours)
  - Column G: Downtime? (Yes/No)
  - Column H: Afterhours? (Yes/No)
  - Metadata row: Project name, ticket number, prepared by, date
  - Downtime explanations, parts, client/vendor dependencies, comments
  - Color coding per task type (remote=blue, onsite=purple, vendor=red,
    client=yellow, procurement=green)

Row layout is auto-detected from the section labels in the Tasklist sheet
("Task - Please separate...", "At completion of project:", "Parts Needed",
"Comments:", etc.). Works on the master template AND on technician copies
whose rows are shifted (e.g. a prior scope copied with the "DO NOT SAVE
OVER" banner row removed). Fails before touching the file if a label is
missing.

Preserves all existing formulas, formatting, and other tabs.
Handles dynamic day counts by inserting rows when tasks exceed the
sheet's available task rows.
Creates automatic backup before writing.

Usage:
  python write_scope.py --input plan.json --file /path/to/scope-copy.xlsx
  cat plan.json | python write_scope.py --file /path/to/scope-copy.xlsx
  python write_scope.py --input plan.json --file scope.xlsx --no-backup
  python write_scope.py --file scope.xlsx --show-layout   # detect only, no write

JSON format:
{
  "metadata": {
    "project_name": "RMCN - New Server",
    "ticket_number": "59542",
    "prepared_by": "Riely Borek",
    "date": "2026-03-18"
  },
  "days": [
    {
      "label": "Day 1 - 8hrs max",
      "tasks": [
        {
          "description": "Rack and cable new server hardware",
          "time_min": 1.5,
          "downtime": "No",
          "afterhours": "No",
          "location": "onsite"
        }
      ]
    }
  ],
  "downtime_explanation": "VM migration requires shutting down VMs...",
  "parts": [
    {
      "description": "Dell PowerEdge R760xs",
      "quantity": 1,
      "part_number": "R760XS-001",
      "url": "https://dell.com/...",
      "price": 5500.00,
      "alternative": "HPE ProLiant DL360"
    }
  ],
  "client_dependencies": "Client must provide downtime approval 48hrs in advance.",
  "cat_herding": "No",
  "vendor_dependencies": {
    "vendors": [
      {"name": "Dell ProSupport"},
      {"name": "ISP - Comcast"}
    ],
    "support_current": "Yes",
    "vendor_charges": "No",
    "existing_ticket": "",
    "contact": "1-800-456-3355",
    "hours": "24/7"
  },
  "comments": "Additional notes..."
}
"""

import argparse
import json
import shutil
import sys
import os
from copy import copy
from datetime import datetime

try:
    import openpyxl
    from openpyxl.styles import PatternFill, Font
except ImportError:
    print("ERROR: openpyxl is required. Install with: pip install openpyxl")
    sys.exit(1)


# --- Section labels used to auto-detect the row layout ---
# key -> (column, lowercase prefix the cell text must start with)
ANCHORS = {
    'meta':          ('A', 'project name'),
    'color_key':     ('A', 'please color code'),
    'task_header':   ('A', 'task - please separate'),
    'at_completion': ('A', 'at completion of project'),
    'update_itglue': ('A', 'update itglue'),
    'deprecate':     ('A', 'deprecate old itglue'),
    'notify':        ('A', 'notify project coordinator'),
    'dt_total':      ('A', 'estimated downtime total'),
    'ah_total':      ('A', 'recommended after hours'),
    'labor':         ('B', 'labor total'),
    'trouble':       ('B', 'troubleshooting'),
    'total':         ('B', 'total (hours)'),
    'pm':            ('B', 'proj mgmt'),
    'dt_question':   ('A', 'if there will be downtime'),
    'parts':         ('A', 'parts needed'),
    'client_dep':    ('A', 'client dependencies'),
    'cat_herding':   ('A', 'would this project involve cat-herding'),
    'vendor_dep':    ('A', 'vendor dependencies'),
    'comments':      ('A', 'comments:'),
}

# Vendor detail labels (column A prefix) -> plan key
VENDOR_DETAIL_LABELS = {
    'is the clients support contract current': 'support_current',
    'does the vendor charge': 'vendor_charges',
    'is there already a ticket open': 'existing_ticket',
    'vendor contact': 'contact',
    'business hours': 'hours',
}

# --- Color coding fills ---
# Template convention:
#   Blue = ITA_office/Remote
#   Purple = ITA_client onsite
#   Red = Vendor Responsibility
#   Yellow = Client Responsibility
#   Green = Procurement
COLOR_MAP = {
    "remote":      PatternFill(start_color="B4C6E7", end_color="B4C6E7", fill_type="solid"),  # Light blue
    "onsite":      PatternFill(start_color="D5A6E6", end_color="D5A6E6", fill_type="solid"),  # Light purple
    "vendor":      PatternFill(start_color="F4CCCC", end_color="F4CCCC", fill_type="solid"),  # Light red
    "client":      PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid"),  # Light yellow
    "procurement": PatternFill(start_color="D9EAD3", end_color="D9EAD3", fill_type="solid"),  # Light green
}

# Day header style
DAY_HEADER_FONT = Font(bold=True, size=11)
DAY_HEADER_FILL = PatternFill(start_color="D9D9D9", end_color="D9D9D9", fill_type="solid")  # Light gray


def _text(ws, coord):
    v = ws[coord].value
    return str(v).strip().lower() if v is not None else ''


def _sum_ranges(col, rows):
    """Compress a sorted row list into SUM ranges, e.g. E51:E54,E56:E58."""
    ranges, start = [], rows[0]
    for a, b in zip(rows, rows[1:] + [None]):
        if b != a + 1:
            ranges.append(f'{col}{start}:{col}{a}' if start != a else f'{col}{a}')
            start = b
    return ','.join(ranges)


def detect_layout(ws):
    """Find section rows by their labels. Returns dict of key -> row(s)."""
    found = {}
    for row in range(1, ws.max_row + 1):
        for key, (col, prefix) in ANCHORS.items():
            if key not in found and _text(ws, f'{col}{row}').startswith(prefix):
                found[key] = row
    missing = [k for k in ANCHORS if k not in found]
    if missing:
        print("ERROR: Could not auto-detect Tasklist layout; missing labels: "
              + ", ".join(f"{ANCHORS[k][0]}='{ANCHORS[k][1]}...'" for k in missing))
        sys.exit(1)

    L = dict(found)
    L['task_start'] = L['task_header'] + 1
    L['task_end'] = L['at_completion'] - 1
    if L['task_end'] < L['task_start']:
        print("ERROR: No task rows between task header and 'At completion of project:'")
        sys.exit(1)

    # Parts: rows between the parts header and Client Dependencies.
    # "Total" labels in column D mark subtotal/grand-total rows; data rows are
    # the non-total rows before the last total row.
    total_rows = [r for r in range(L['parts'] + 1, L['client_dep'])
                  if _text(ws, f'D{r}').startswith('total')]
    last = total_rows[-1] if total_rows else L['client_dep']
    L['parts_total_rows'] = total_rows
    L['parts_rows'] = [r for r in range(L['parts'] + 1, last) if r not in total_rows]

    # Vendor detail rows (label-matched between vendor header and comments)
    L['vendor_rows'] = {}
    for r in range(L['vendor_dep'] + 1, L['comments']):
        t = _text(ws, f'A{r}')
        for prefix, key in VENDOR_DETAIL_LABELS.items():
            if t.startswith(prefix):
                L['vendor_rows'][key] = r
    L['vendor_name_row'] = L['vendor_dep'] + 1

    # Metadata style. Master template: "Project Name:" row holds label+value
    # strings in A/B/C. Copied scopes may keep labels in that row and hold the
    # ticket / prepared-by VALUES in B/C of the row below (the color-key row).
    below = L['meta'] + 1
    L['meta_values_below'] = (
        below == L['color_key']
        and (ws[f'B{below}'].value not in (None, '') or ws[f'C{below}'].value not in (None, ''))
    )
    return L


def describe_layout(L):
    parts = L['parts_rows']
    return (f"  Layout: meta row {L['meta']}"
            f"{' (values in row ' + str(L['meta'] + 1) + ')' if L['meta_values_below'] else ''}, "
            f"task rows {L['task_start']}-{L['task_end']}, "
            f"at-completion {L['at_completion']}, "
            f"downtime explanation A{L['dt_question'] + 1}, "
            f"parts rows {parts[0] if parts else '-'}-{parts[-1] if parts else '-'} "
            f"(total rows {L['parts_total_rows']}), "
            f"client deps {L['client_dep']}, vendor deps {L['vendor_dep']}, comments {L['comments']}")


def load_plan(input_path=None):
    """Load task plan from file or stdin."""
    if input_path:
        with open(input_path, 'r') as f:
            return json.load(f)
    else:
        return json.load(sys.stdin)


def count_task_rows(plan):
    """Count total rows needed for all days (day headers + tasks)."""
    total = 0
    for day in plan.get('days', []):
        total += 1  # Day header row
        total += len(day.get('tasks', []))
    return total


def copy_row_style(ws, source_row, target_row, max_col=9):
    """Copy cell formatting from source_row to target_row."""
    for col in range(1, max_col + 1):
        src_cell = ws.cell(row=source_row, column=col)
        tgt_cell = ws.cell(row=target_row, column=col)
        if src_cell.has_style:
            tgt_cell.font = copy(src_cell.font)
            tgt_cell.border = copy(src_cell.border)
            tgt_cell.fill = copy(src_cell.fill)
            tgt_cell.number_format = src_cell.number_format
            tgt_cell.alignment = copy(src_cell.alignment)
            tgt_cell.protection = copy(src_cell.protection)


def apply_task_color(ws, row, location, max_col=9):
    """Apply color fill to Column A only for a task row based on location type."""
    if not location:
        return
    fill = COLOR_MAP.get(location.lower())
    if fill:
        ws.cell(row=row, column=1).fill = fill  # Column A only


def apply_day_header_style(ws, row, max_col=9):
    """Apply bold gray style to day header rows."""
    for col in range(1, max_col + 1):
        cell = ws.cell(row=row, column=col)
        cell.font = DAY_HEADER_FONT
        cell.fill = DAY_HEADER_FILL


def create_backup(xlsx_path):
    """Create a timestamped backup of the file before modifying."""
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    base, ext = os.path.splitext(xlsx_path)
    backup_path = f"{base}_backup_{timestamp}{ext}"
    shutil.copy2(xlsx_path, backup_path)
    print(f"  Backup created: {backup_path}")
    return backup_path


def write_scope(plan, xlsx_path, no_backup=False):
    """Write the task plan into the Tasklist tab of the scope workbook."""
    if not os.path.exists(xlsx_path):
        print(f"ERROR: File not found: {xlsx_path}")
        sys.exit(1)

    wb = openpyxl.load_workbook(xlsx_path)

    if 'Tasklist' not in wb.sheetnames:
        print("ERROR: 'Tasklist' sheet not found in workbook")
        sys.exit(1)

    ws = wb['Tasklist']

    # --- Detect layout and validate before backup/write ---
    L = detect_layout(ws)
    print(describe_layout(L))

    parts = plan.get('parts', [])
    if len(parts) > len(L['parts_rows']):
        print(f"ERROR: {len(parts)} parts but only {len(L['parts_rows'])} part rows available "
              f"(rows {L['parts_rows']}); add rows in the workbook or trim parts")
        sys.exit(1)

    # --- Create backup ---
    if not no_backup:
        create_backup(xlsx_path)

    # --- Calculate row needs ---
    rows_needed = count_task_rows(plan)
    rows_available = L['task_end'] - L['task_start'] + 1
    extra_rows = max(0, rows_needed - rows_available)

    # Capture style from a representative task row before insertion
    style_source_row = L['task_start'] + 1

    # Insert extra rows if needed, then re-detect (everything below shifted)
    if extra_rows > 0:
        ws.insert_rows(L['task_end'] + 1, amount=extra_rows)
        for i in range(extra_rows):
            copy_row_style(ws, style_source_row, L['task_end'] + 1 + i)
        print(f"  Inserted {extra_rows} extra rows to accommodate {rows_needed} task rows")
        L = detect_layout(ws)

    task_start, task_end = L['task_start'], L['task_end']

    # --- Write metadata ---
    meta = plan.get('metadata', {})
    mr = L['meta']
    if meta.get('project_name'):
        ws[f'A{mr}'] = f"Project Name: {meta['project_name']}"
    if L['meta_values_below']:
        # Copied-scope style: labels stay in row mr, values go in row mr+1
        vr = mr + 1
        if meta.get('ticket_number'):
            t = str(meta['ticket_number'])
            ws[f'B{vr}'] = int(t) if t.isdigit() else t
        if meta.get('prepared_by'):
            ws[f'C{vr}'] = meta['prepared_by']
    else:
        # Master-template style: combined label + value in the Project Name row
        if meta.get('ticket_number'):
            ws[f'B{mr}'] = f"Project Ticket Number: {meta['ticket_number']}"
        if meta.get('prepared_by'):
            ws[f'C{mr}'] = f"Sales Scope Prepared by: {meta['prepared_by']}"
    if meta.get('date'):
        ws[f'G{mr}'] = meta['date']

    # --- Clear existing task content (values, fills, stale bold) ---
    for row in range(task_start, task_end + 1):
        for col in ['A', 'B', 'C', 'G', 'H', 'I', 'J']:
            ws[f'{col}{row}'] = None
        for col in range(1, 10):
            ws.cell(row=row, column=col).fill = PatternFill(fill_type=None)
        a = ws.cell(row=row, column=1)
        if a.font is not None and a.font.b:
            f = copy(a.font)
            f.b = False
            a.font = f

    # --- Write task rows ---
    days = plan.get('days', [])
    current_row = task_start
    day_hours = {}

    for day_idx, day in enumerate(days):
        # Write day header
        day_label = day.get('label', f'Day {day_idx + 1} - 8hrs max')
        ws[f'A{current_row}'] = day_label
        ws[f'D{current_row}'] = f'=(C{current_row}+E{current_row})/2'
        ws[f'E{current_row}'] = f'=C{current_row}*2'
        apply_day_header_style(ws, current_row)
        current_row += 1

        # Write tasks for this day
        tasks = day.get('tasks', [])
        total_hours = 0
        for task in tasks:
            ws[f'A{current_row}'] = task.get('description', '')
            time_min = task.get('time_min')
            if time_min is not None:
                ws[f'C{current_row}'] = time_min
                total_hours += time_min
            if task.get('downtime'):
                ws[f'G{current_row}'] = task['downtime']
            if task.get('afterhours'):
                ws[f'H{current_row}'] = task['afterhours']
            ws[f'D{current_row}'] = f'=(C{current_row}+E{current_row})/2'
            ws[f'E{current_row}'] = f'=C{current_row}*2'

            # Apply color coding
            location = task.get('location', '')
            if location:
                apply_task_color(ws, current_row, location)

            current_row += 1

        day_hours[day_label] = total_hours
        if total_hours > 8:
            print(f"  \u26a0\ufe0f  WARNING: {day_label} totals {total_hours}hrs (exceeds 8hr max)")

    # --- "At completion" rows (labels already present; refresh formulas) ---
    for key in ('update_itglue', 'deprecate'):
        r = L[key]
        ws[f'D{r}'] = f'=(C{r}+E{r})/2'
        ws[f'E{r}'] = f'=C{r}*2'

    # --- Summary formulas against detected ranges ---
    completion_last_row = L['notify']

    r = L['dt_total']
    ws[f'C{r}'] = f'=SUMIF(G{task_start}:G{task_end},"Yes",C{task_start}:C{task_end})'
    ws[f'D{r}'] = f'=(C{r}+E{r})/2'
    ws[f'E{r}'] = f'=C{r}*2'

    r = L['ah_total']
    ws[f'C{r}'] = f'=SUMIF(H{task_start}:H{task_end},"Yes",C{task_start}:C{task_end})'
    ws[f'D{r}'] = f'=(C{r}+E{r})/2'
    ws[f'E{r}'] = f'=C{r}*2'

    lr, tr, tot, pm = L['labor'], L['trouble'], L['total'], L['pm']
    for col in 'CDE':
        ws[f'{col}{lr}'] = f'=SUM({col}{task_start}:{col}{completion_last_row})'
        ws[f'{col}{tr}'] = f'=SUM({col}{lr}*0.25)'
        ws[f'{col}{tot}'] = f'=SUM({col}{lr}+{col}{tr})'
        ws[f'{col}{pm}'] = f'=SUM({col}{task_start}:{col}{task_end})*0.2'

    # --- Downtime Explanation (row after the "If there will be downtime" prompt) ---
    dt_explanation = plan.get('downtime_explanation', '')
    ws[f'A{L["dt_question"] + 1}'] = dt_explanation if dt_explanation else "Please be specific"

    # --- Parts Needed ---
    for r in L['parts_rows']:
        for col in 'ABCDF':
            ws[f'{col}{r}'] = None
        ws[f'E{r}'] = 0
    for part, r in zip(parts, L['parts_rows']):
        ws[f'A{r}'] = part.get('description', '')
        if part.get('quantity') is not None:
            ws[f'B{r}'] = part['quantity']
        if part.get('part_number'):
            ws[f'C{r}'] = part['part_number']
        if part.get('url'):
            ws[f'D{r}'] = part['url']
        if part.get('price') is not None:
            ws[f'E{r}'] = part['price']
        if part.get('alternative'):
            ws[f'F{r}'] = part['alternative']

    # Total rows: a subtotal sums data rows since the previous total row; the
    # last (grand) total sums ALL data rows and never includes subtotal rows.
    prev = L['parts']
    for i, total_row in enumerate(L['parts_total_rows']):
        is_last = i == len(L['parts_total_rows']) - 1
        rows = L['parts_rows'] if is_last else [r for r in L['parts_rows'] if prev < r < total_row]
        if rows:
            ws[f'E{total_row}'] = f'=SUM({_sum_ranges("E", rows)})'
        prev = total_row

    # --- Client Dependencies (text goes below the cat-herding question) ---
    client_deps = plan.get('client_dependencies', '')
    if client_deps:
        ws[f'A{L["cat_herding"] + 1}'] = client_deps
    cat_herding = plan.get('cat_herding', '')
    if cat_herding:
        ws[f'B{L["cat_herding"]}'] = cat_herding

    # --- Vendor Dependencies ---
    vendor_deps = plan.get('vendor_dependencies', {})
    for i, v in enumerate(vendor_deps.get('vendors', [])[:3]):
        col = ['B', 'C', 'D'][i]
        ws[f'{col}{L["vendor_name_row"]}'] = f"Vendor {i+1}: {v.get('name', '')}"
    for key, r in L['vendor_rows'].items():
        if vendor_deps.get(key):
            ws[f'B{r}'] = vendor_deps[key]

    # --- Comments ---
    comments = plan.get('comments', '')
    if comments:
        ws[f'A{L["comments"] + 1}'] = comments

    # --- Save ---
    wb.save(xlsx_path)
    print(f"SUCCESS: Task plan written to {xlsx_path}")
    print(f"  Days: {len(days)}")
    total_tasks = sum(len(d.get('tasks', [])) for d in days)
    print(f"  Tasks: {total_tasks}")
    print(f"  Task rows: {task_start}-{current_row - 1}")
    if extra_rows > 0:
        print(f"  Extra rows inserted: {extra_rows}")
    if parts:
        print(f"  Parts: {len(parts)}")
    for label, hrs in day_hours.items():
        status = "\u2705" if hrs <= 8 else "\u26a0\ufe0f  OVER"
        print(f"  {label}: {hrs}hrs {status}")


def main():
    parser = argparse.ArgumentParser(
        description="Write scope task plan to Excel template"
    )
    parser.add_argument(
        '--input', '-i',
        help="Path to JSON plan file (default: stdin)"
    )
    parser.add_argument(
        '--file', '-f', required=True,
        help="Path to the Excel scope template copy"
    )
    parser.add_argument(
        '--no-backup', action='store_true',
        help="Skip creating a backup before writing"
    )
    parser.add_argument(
        '--show-layout', action='store_true',
        help="Only print the detected Tasklist layout; do not write"
    )
    args = parser.parse_args()

    if args.show_layout:
        if not os.path.exists(args.file):
            print(f"ERROR: File not found: {args.file}")
            sys.exit(1)
        ws = openpyxl.load_workbook(args.file)['Tasklist']
        print(describe_layout(detect_layout(ws)))
        return

    plan = load_plan(args.input)
    write_scope(plan, args.file, no_backup=args.no_backup)


if __name__ == '__main__':
    main()
